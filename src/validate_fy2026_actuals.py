from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from infosys_valuation_model import DATA, OUTPUTS, forecast_three_statement, load_inputs


def format_value(value, unit):
    if pd.isna(value):
        return ""
    if unit == "percentage":
        return f"{value:.1%}"
    return f"{value:,.0f}"


def build_validation():
    historical, _, assumptions = load_inputs()
    forecast = forecast_three_statement(historical, assumptions)
    fy26_forecast = forecast.loc[forecast["Fiscal Year"] == "FY2026"].iloc[0]
    actuals = pd.read_csv(DATA / "fy2026_actuals.csv").set_index("Metric")

    forecast_map = {
        "Revenue": fy26_forecast["Revenue"],
        "Revenue Growth": fy26_forecast["Revenue Growth"],
        "Operating Profit": fy26_forecast["EBIT"],
        "Operating Margin": fy26_forecast["EBIT Margin"],
        "Adjusted Operating Profit": fy26_forecast["EBIT"],
        "Adjusted Operating Margin": fy26_forecast["EBIT Margin"],
        "CFO": fy26_forecast["CFO"],
        "Company FCF": fy26_forecast["FCF"],
        "D&A": fy26_forecast["D&A"],
        "Capex": fy26_forecast["Capex"],
    }

    rows = []
    for metric, forecast_value in forecast_map.items():
        actual_value = float(actuals.loc[metric, "Actual"])
        unit = actuals.loc[metric, "Unit"]
        variance = actual_value - forecast_value
        pct_error = variance / forecast_value if forecast_value else None
        rows.append(
            {
                "Metric": metric,
                "Forecast": forecast_value,
                "Actual FY2026": actual_value,
                "Variance": variance,
                "Percent Error": pct_error,
                "Unit": unit,
                "Source": actuals.loc[metric, "Source"],
                "Notes": actuals.loc[metric, "Notes"],
            }
        )

    validation = pd.DataFrame(rows)
    validation["Forecast Display"] = validation.apply(
        lambda row: format_value(row["Forecast"], row["Unit"]), axis=1
    )
    validation["Actual Display"] = validation.apply(
        lambda row: format_value(row["Actual FY2026"], row["Unit"]), axis=1
    )
    validation["Variance Display"] = validation.apply(
        lambda row: (
            f"{row['Variance']:.1%}"
            if row["Unit"] == "percentage"
            else f"{row['Variance']:,.0f}"
        ),
        axis=1,
    )
    validation["Percent Error Display"] = validation["Percent Error"].map(lambda x: f"{x:.1%}")
    return actuals.reset_index(), validation


def make_chart(validation):
    OUTPUTS.mkdir(exist_ok=True)
    chart_data = validation[validation["Metric"].isin(["Revenue", "Operating Profit", "CFO", "Company FCF"])].copy()
    x = range(len(chart_data))
    width = 0.36

    fig, ax = plt.subplots(figsize=(8, 4.4))
    ax.bar([i - width / 2 for i in x], chart_data["Forecast"], width, label="Forecast", color="#6b7c8d")
    ax.bar([i + width / 2 for i in x], chart_data["Actual FY2026"], width, label="Actual FY2026", color="#007cc3")
    ax.set_xticks(list(x))
    ax.set_xticklabels(chart_data["Metric"], rotation=0)
    ax.set_ylabel("INR crore")
    ax.set_title("FY2026 Forecast vs Actual")
    ax.legend(frameon=False)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout()

    path = OUTPUTS / "fy2026_forecast_vs_actual.png"
    fig.savefig(path, dpi=160)
    plt.close(fig)
    return path


def write_excel(actuals, validation):
    workbook_path = OUTPUTS / "fy2026_prediction_validation.xlsx"
    display_cols = [
        "Metric",
        "Forecast Display",
        "Actual Display",
        "Variance Display",
        "Percent Error Display",
        "Notes",
        "Source",
    ]
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        actuals.to_excel(writer, sheet_name="FY2026 Actuals", index=False)
        validation.to_excel(writer, sheet_name="Validation Raw", index=False)
        validation[display_cols].to_excel(writer, sheet_name="Validation Summary", index=False)

    wb = load_workbook(workbook_path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
        for col in range(1, ws.max_column + 1):
            letter = get_column_letter(col)
            width = max(len(str(ws.cell(row=row, column=col).value or "")) for row in range(1, ws.max_row + 1))
            ws.column_dimensions[letter].width = min(width + 2, 42)
    wb.save(workbook_path)
    return workbook_path


def write_memo(validation, chart_path):
    def row(metric):
        return validation.loc[validation["Metric"] == metric].iloc[0]

    revenue = row("Revenue")
    op_margin = row("Operating Margin")
    adj_margin = row("Adjusted Operating Margin")
    fcf = row("Company FCF")
    cfo = row("CFO")

    memo = f"""# FY2026 Prediction Validation

## Verdict

The FY2026 forecast was too conservative on revenue and cash generation. Margin quality was closer: reported operating margin missed the forecast because of the Labour Codes provision, but adjusted operating margin was slightly ahead of the model.

## Forecast vs Actual

- Revenue: forecast {revenue['Forecast Display']} vs actual {revenue['Actual Display']} INR crore, a {revenue['Percent Error Display']} beat.
- Operating margin, reported: forecast {op_margin['Forecast Display']} vs actual {op_margin['Actual Display']}, a {op_margin['Variance'] * 10000:,.0f} bps miss.
- Operating margin, adjusted: forecast {adj_margin['Forecast Display']} vs actual {adj_margin['Actual Display']}, a {adj_margin['Variance'] * 10000:,.0f} bps beat.
- CFO: forecast {cfo['Forecast Display']} vs actual {cfo['Actual Display']} INR crore, a {cfo['Percent Error Display']} beat.
- FCF: forecast {fcf['Forecast Display']} vs company-reported actual {fcf['Actual Display']} INR crore, a {fcf['Percent Error Display']} beat.

## Finance Read

- Growth was the main miss: the model assumed 3.0% INR revenue growth, while Infosys delivered 9.6%.
- Reported operating profit was affected by a one-off Labour Codes provision of INR 1,289 crore.
- Adjusted operating margin of 21.0% supports the original margin thesis better than the reported 20.3% figure.
- Cash generation came in stronger than forecast, so the DCF downside was probably too punitive if FY2026 becomes the new base year.

## Source Notes

- Infosys FY2026 data sheet: https://www.infosys.com/investors/reports-filings/financials/data-sheet.html
- Infosys Q4 FY2026 IFRS INR press release: https://www.infosys.com/investors/reports-filings/quarterly-results/2025-2026/q4/documents/ifrs-inr-press-release.pdf
- Infosys FY2026 consolidated financial statements: https://www.infosys.com/investors/reports-filings/quarterly-results/2025-2026/q4/documents/consolidated/consol-fy26-annual-finstatement.pdf

![FY2026 forecast vs actual]({chart_path.as_posix()})
"""
    path = OUTPUTS / "fy2026_validation_memo.md"
    path.write_text(memo, encoding="utf-8")
    return path


def main():
    actuals, validation = build_validation()
    OUTPUTS.mkdir(exist_ok=True)
    validation.to_csv(OUTPUTS / "fy2026_prediction_validation.csv", index=False)
    chart_path = make_chart(validation)
    workbook_path = write_excel(actuals, validation)
    memo_path = write_memo(validation, chart_path)
    print(f"Built {OUTPUTS / 'fy2026_prediction_validation.csv'}")
    print(f"Built {workbook_path}")
    print(f"Built {memo_path}")
    print(f"Built {chart_path}")


if __name__ == "__main__":
    main()
